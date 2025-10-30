import PyPDF2
from docx import Document
import nltk
from io import BytesIO

# Download required NLTK data
try:
    nltk.data.find('tokenizers/punkt')
except LookupError:
    nltk.download('punkt')
    
try:
    nltk.data.find('tokenizers/punkt_tab')
except LookupError:
    nltk.download('punkt_tab')

def extract_text_from_pdf(file_bytes):
    """Extract text from PDF file"""
    pdf_reader = PyPDF2.PdfReader(BytesIO(file_bytes))
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text() + "\n"
    return text

def extract_text_from_docx(file_bytes):
    """Extract text from DOCX file"""
    doc = Document(BytesIO(file_bytes))
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def extract_text_from_txt(file_bytes):
    """Extract text from TXT file"""
    return file_bytes.decode('utf-8')

def segment_text(text, max_chars=400):
    """Split text into segments for translation, preserving placeholders"""
    import re
    
    # Extract placeholders and replace with markers
    placeholders = {}
    placeholder_pattern = r'\[(TABLE_\d{3}|FIGURE_\d{3})\]'
    
    def replace_placeholder(match):
        marker = f"__PLACEHOLDER_{len(placeholders)}__"
        placeholders[marker] = match.group(0)
        return marker
    
    text_with_markers = re.sub(placeholder_pattern, replace_placeholder, text)
    
    # Segment the text
    try:
        sentences = nltk.sent_tokenize(text_with_markers)
    except LookupError:
        # Fallback to simple splitting if NLTK fails
        sentences = text_with_markers.split('. ')
        sentences = [s + '.' for s in sentences[:-1]] + [sentences[-1]]
    
    segments = []
    current_segment = ""
    
    for sentence in sentences:
        # Check if sentence contains a placeholder marker
        if '__PLACEHOLDER_' in sentence:
            # Add current segment if exists
            if current_segment:
                segments.append(current_segment.strip())
                current_segment = ""
            # Add placeholder as separate segment
            segments.append(sentence.strip())
        elif len(current_segment + sentence) <= max_chars:
            current_segment += sentence + " "
        else:
            if current_segment:
                segments.append(current_segment.strip())
            current_segment = sentence + " "
    
    if current_segment:
        segments.append(current_segment.strip())
    
    # Restore placeholders
    restored_segments = []
    for segment in segments:
        for marker, placeholder in placeholders.items():
            segment = segment.replace(marker, placeholder)
        restored_segments.append(segment)
    
    return restored_segments

def process_document(file):
    """Process uploaded document and extract text"""
    # Handle Gradio file object
    if hasattr(file, 'name'):
        file_path = file.name
    else:
        file_path = file
    
    file_extension = file_path.lower().split('.')[-1]
    
    if file_extension == 'pdf':
        # Read file content for PDF
        with open(file_path, 'rb') as f:
            file_bytes = f.read()
        text = extract_text_from_pdf(file_bytes)
        return text, None  # No element processor for PDF yet
    elif file_extension == 'docx':
        # Use enhanced processor for DOCX
        from document_elements import DocumentProcessor
        processor = DocumentProcessor()
        text = processor.extract_docx_elements(file_path)
        return text, processor
    elif file_extension == 'txt':
        # Read file content for TXT
        with open(file_path, 'rb') as f:
            file_bytes = f.read()
        text = extract_text_from_txt(file_bytes)
        return text, None  # No element processor for TXT
    else:
        raise ValueError(f"Unsupported file format: {file_extension}")

def process_document_simple(file):
    """Simple document processing (backward compatibility)"""
    text, _ = process_document(file)
    return text